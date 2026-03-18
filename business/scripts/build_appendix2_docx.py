from pathlib import Path
from shutil import copy2

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
OFFICIAL = ROOT / "official_docs"
GENERATED = ROOT / "generated"
TEMPLATE = next(
    p for p in OFFICIAL.iterdir() if p.suffix.lower() == ".docx" and "증빙서류 제출목록 안내" in p.name
)
OUTPUT = GENERATED / "사업계획서_PULSE_2026_별첨2_증빙서류_입력본.docx"


def build() -> Path:
    GENERATED.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        OUTPUT.unlink()
    copy2(TEMPLATE, OUTPUT)
    doc = Document(str(OUTPUT))
    tables = doc.tables

    # Table 2: submission checklist
    checklist = tables[1]
    checklist.cell(1, 2).text = (
        "(필수) 대표자(신청자) 신분증 사본 [붙임 1]\n"
        "- 주민등록증, 운전면허증, 여권(구형) 중 1개\n"
        "- 본인 확인용 제출"
    )
    checklist.cell(2, 2).text = ""
    checklist.cell(3, 2).text = ""
    checklist.cell(4, 1).text = ""
    checklist.cell(4, 2).text = ""
    checklist.cell(5, 1).text = ""
    checklist.cell(5, 2).text = ""
    checklist.cell(6, 1).text = ""
    checklist.cell(6, 2).text = ""
    checklist.cell(7, 1).text = "[붙임 7]"
    checklist.cell(7, 2).text = (
        "(해당 시) 제출한 자료 외 사업계획서 평가 참고용 증빙 [붙임 7]\n"
        "1. KCI 논문 등재 증빙\n"
        "2. MVP 상세 화면 캡처\n"
        "3. 프로젝트 기술 아키텍처\n"
        "4. 사용자 흐름도"
    )

    # Table 3: ID copy page
    id_page = tables[2]
    id_page.cell(1, 0).text = (
        "[붙임 1] 대표자(신청자) 신분증 사본 첨부 위치\n\n"
        "- 주민등록증, 운전면허증, 여권(구형) 중 1개 첨부\n"
        "- 주민등록번호 뒷자리는 마스킹 후 제출\n"
        "- 컬러 또는 흑백 스캔본 모두 가능"
    )

    # Tables 4-8: mark as not applicable
    for idx in range(3, 8):
        tables[idx].cell(1, 0).text = "해당 없음"

    # Table 9: appendix evidence page
    appendix = tables[8]
    appendix.cell(1, 0).text = (
        "[붙임 7] 기타 증빙서류 제출 순서\n\n"
        "1. [붙임 7-1] KCI 논문 등재 증빙\n"
        "   - KCI 검색 결과 화면 또는 논문 첫 페이지 캡처 삽입\n"
        "   - 사업계획서 본문 내 문제 정의 및 신뢰도 보강 근거와 연결\n\n"
        "2. [붙임 7-2] MVP 상세 화면 캡처\n"
        "   - 리뷰 분석 및 페르소나 도출 화면 캡처 삽입\n"
        "   - 홍보영상 생성 화면 캡처 삽입\n\n"
        "3. [붙임 7-3] 프로젝트 전체 기술 아키텍처\n"
        "   - 프론트엔드, 메인 백엔드, AI 서버 구조가 보이도록 삽입\n\n"
        "4. [붙임 7-4] 사용자 흐름도\n"
        "   - 고객 이해 → 홍보 제작 → 실행 제안 흐름이 보이도록 삽입\n\n"
        "※ 권장 편집 순서: 신분증 사본 뒤에 [붙임 7-1]부터 [붙임 7-4] 순으로 배치\n"
        "※ 각 자료 상단에 [붙임 7-1], [붙임 7-2] 등 표기를 유지하면 검토가 쉽습니다."
    )

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
